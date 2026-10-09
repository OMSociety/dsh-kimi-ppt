#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pptd_to_pptx.py — 本地 .pptd -> .pptx 转换器（纯本地版）
不依赖无头浏览器 / agent-browser / www.kimi.com，只用 python-pptx + 本地文件。

用法:
  python pptd_to_pptx.py <deck.pptd> [-o <out.pptx>]
    -o 缺省：输出到项目目录下同名 .pptx（不覆盖已有文件，除非 --force）

支持的元素: text / shape(常见几何形) / line / image / table；icon 与 chart 静默跳过，image 只插本地存在的文件，未知 shapeName 退化为矩形。
"""
import os, sys, re, argparse, html
import yaml
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

PX = 12700          # 1px = 1pt = 12700 EMU（规范：1px == 1pt）

# 规范名 -> 实装名（与 reference/local-fonts.md 一致；本地导出写实装名让 PowerPoint 命中本机字体）
CANON_TO_INSTALLED = {
    "MiSans": "MiSans", "Noto Sans SC": "Noto Sans SC", "思源宋体": "思源宋体 CN",
    "Source Han Serif": "思源宋体 CN", "阿里妈妈刀隶体": "阿里妈妈刀隶体",
    "阿里妈妈东方大楷": "阿里妈妈东方大楷", "阿里妈妈数黑体": "阿里妈妈数黑体",
    "站酷文艺体": "站酷文艺体", "ZCOOL KuaiLe": "ZCOOL KuaiLe",
    "得意黑": "得意黑 斜体", "Smiley Sans": "得意黑 斜体", "飞波正点体": "飞波正点体",
    "霞鹜新致宋": "霞鹜新致宋＋", "LXGW Bright": "霞鹜文楷", "霞鹜文楷": "霞鹜文楷",
    "精品点阵体": "精品点阵体9×9 1.93 R", "Liter": "Liter",
    "HedvigLettersSans": "Hedvig Letters Sans", "Oranienbaum": "Oranienbaum",
    "QuattrocentoSans": "Quattrocento Sans", "Unna": "Unna", "Coda": "Coda",
    "Jersey15": "Jersey 15", "Jersey20Charted": "Jersey 20 Charted",
    "SortsMillGoudy": "Sorts Mill Goudy", "更纱黑体 SC": "更纱黑体 SC",
    "Sarasa Gothic": "更纱黑体 SC", "Microsoft YaHei": "Microsoft YaHei",
    "微软雅黑": "Microsoft YaHei", "SimHei": "SimHei", "黑体": "SimHei",
    "SimSun": "SimSun", "宋体": "SimSun", "FangSong": "FangSong", "仿宋": "FangSong",
    "KaiTi": "KaiTi", "楷体": "KaiTi", "思源黑体 CN": "思源黑体 CN",
    "Century Gothic": "Century Gothic", "Bahnschrift": "Bahnschrift",
    "Georgia": "Georgia", "Cambria": "Cambria", "Constantia": "Constantia",
}

def set_rpr_fonts(run, latin=None, ea=None):
    """同时设置 latin 与 East Asian 两种 typeface，避免中文落回默认字体。规范名先映射为实装名。"""
    latin = CANON_TO_INSTALLED.get(latin, latin) if latin else latin
    ea = CANON_TO_INSTALLED.get(ea, ea) if ea else ea
    if latin:
        run.font.name = latin                        # 处理 <a:latin>（python-pptx 会放对位置）
    if ea:
        rPr = run._r.get_or_add_rPr()
        ea_el = rPr.find(qn('a:ea'))
        if ea_el is None:
            latin_el = rPr.find(qn('a:latin'))
            ea_el = rPr.makeelement(qn('a:ea'), {})
            if latin_el is not None:
                latin_el.addnext(ea_el)
            else:
                rPr.append(ea_el)
        ea_el.set('typeface', ea)

# ---- 常见 shapeName -> MSO_SHAPE 属性名（惰性解析，缺省退回矩形） ----
SHAPE_MAP = {
    "rect": "RECTANGLE",
    "roundRect": "ROUNDED_RECTANGLE",
    "ellipse": "OVAL",
    "circle": "OVAL",
    "triangle": "ISOSCELES_TRIANGLE",
    "rtTriangle": "RIGHT_TRIANGLE",
    "diamond": "DIAMOND",
    "chevron": "CHEVRON",
    "homePlate": "PENTAGON",
    "pentagon": "PENTAGON",
    "donut": "DONUT",
    "star5": "STAR_5_POINT",
    "rightArrow": "RIGHT_ARROW",
    "leftArrow": "LEFT_ARROW",
    "bracePair": "BRACE_PAIR",
    "hexagon": "HEXAGON",
    "blockArc": "BLOCK_ARC",
    "trapezoid": "TRAPEZOID",
}

def find_shape(name):
    """按 shapeName 找 MSO_SHAPE；找不到返回 RECTANGLE。"""
    attr = SHAPE_MAP.get(name)
    if attr:
        got = getattr(MSO_SHAPE, attr, None)
        if got is not None:
            return got
    return MSO_SHAPE.RECTANGLE

# ---- 静默降级清单：凡未按 spec 渲染的元素/字段都记一条，结束统一打印 ----
DROPPED = []
_NOTE_SEEN = set()

def _note(element_type, element_id, reason):
    key = (str(element_type), str(element_id), reason)
    if key in _NOTE_SEEN:
        return
    _NOTE_SEEN.add(key)
    DROPPED.append(key)

def _crop_fractions(crop):
    """ImageCrop 四边比例（缺省 0）。负值（outset）或退化源区返回 None，由调用方降级。"""
    if not crop:
        return 0.0, 0.0, 0.0, 0.0
    l = float(crop.get("left", 0) or 0)
    t = float(crop.get("top", 0) or 0)
    r = float(crop.get("right", 0) or 0)
    b = float(crop.get("bottom", 0) or 0)
    if min(l, t, r, b) < 0 or l + r >= 1 or t + b >= 1:
        return None
    return l, t, r, b

def _fit_rect(iw, ih, w, h, l, t, r, b, mode):
    """按渲染链 crop→fit 算源图裁剪比例（全图比例）与显示框。
    cover：显示框=bounds，多出部分进 srcRect 裁掉；contain：整幅子图等比放入框内居中；
    fill：直接拉伸（仅保留 crop 的 srcRect）。"""
    sw = 1.0 - l - r
    sh = 1.0 - t - b
    src_ar = (iw * sw) / (ih * sh)
    box_ar = w / h
    if mode == "contain":
        if src_ar >= box_ar:
            return (l, t, r, b), w, w / src_ar
        return (l, t, r, b), h * src_ar, h
    if mode == "cover":
        if src_ar > box_ar:
            extra = (1.0 - box_ar / src_ar) / 2.0
            return (l + sw * extra, t, r + sw * extra, b), w, h
        if src_ar < box_ar:
            extra = (1.0 - src_ar / box_ar) / 2.0
            return (l, t + sh * extra, r, b + sh * extra), w, h
    return (l, t, r, b), w, h

def _set_src_rect(pic, l, t, r, b):
    """把源图裁剪比例写进 blipFill 的 a:srcRect（ST_Percentage = 比例 × 100000）。"""
    vals = {k: int(round(v * 100000)) for k, v in (("l", l), ("t", t), ("r", r), ("b", b))}
    if not any(vals.values()):
        return
    blip = pic._element.blipFill.blip
    src_rect = pic._element.blipFill.makeelement(
        qn("a:srcRect"), {k: str(v) for k, v in vals.items() if v})
    blip.addnext(src_rect)  # blipFill 子序：blip → srcRect → stretch

COLORS = {
    "red": "E30613", "yellow": "FFC300", "blue": "0066B3", "black": "111111",
    "white": "FFFFFF", "gray": "F4F4F2", "text": "222222", "accent": "FFC300",
    "primary": "0066B3", "secondary": "555555", "success": "1D8348",
    "warning": "B7950B", "danger": "C0392B",
}
# 若主题里有 undefined，则按常见颜色名兜底；其它走 hex。

def resolve_color(c, theme_colors, fallback=None):
    if c is None:
        return fallback if fallback else "000000"
    if isinstance(c, str):
        s = c.strip()
        if s.startswith("$"):
            key = s[1:]
            if key in theme_colors:
                return theme_colors[key].replace("#", "")
            if key in COLORS:
                return COLORS[key]
            return fallback if fallback else "000000"
        return s.replace("#", "")
    return "000000"

def rgb(c):
    c = c.lstrip("#")
    if len(c) == 8:
        c = c[:6]
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    return RGBColor.from_string(c or "000000")

def to_emu(px):
    return Emu(int(round(float(px) * PX)))

def apply_fill(shape, fill, theme_colors, default, owner=None):
    if fill is None:
        shape.fill.background()          # pptd.md: default "not applied" —— 无填充
        return
    typ = fill.get("type", "solid")
    if typ == "solid":
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(resolve_color(fill.get("color"), theme_colors))
    elif typ == "gradient":
        stops = fill.get("stops") or []
        try:
            shape.fill.gradient()
            gs = shape.fill.gradient_stops
            if len(gs) >= 1 and stops:
                gs[0].color.rgb = rgb(resolve_color(stops[0].get("color"), theme_colors))
            if len(gs) >= 2 and len(stops) >= 2:
                gs[1].color.rgb = rgb(resolve_color(stops[-1].get("color"), theme_colors))
            ang = fill.get("angle", 90)
            try:
                shape.fill.gradient_angle = float(ang)
            except Exception:
                pass
        except Exception:
            shape.fill.solid()
            if stops:
                shape.fill.fore_color.rgb = rgb(resolve_color(stops[0].get("color"), theme_colors))
            else:
                shape.fill.fore_color.rgb = default
    elif typ == "image":
        shape.fill.background()
        if owner:
            _note(owner[0], owner[1], "image fill on shape not supported in local track")

def apply_border(shape, border, theme_colors):
    if not border:
        return
    shape.line.color.rgb = rgb(resolve_color(border.get("color"), theme_colors, "000000"))
    shape.line.width = Emu(int(round(float(border.get("width", 1)) * PX)))

_UNSUPPORTED_CELL_FIELDS = ("lineHeight", "lineHeightPx", "letterSpacing", "marginTop", "backgroundColor")

def _first(*vals):
    for v in vals:
        if v is not None:
            return v
    return None

def _resolve_table_style(el, theme):
    st = el.get("style")
    if isinstance(st, str) and st.startswith("$"):
        return (theme.get("tableStyles") or {}).get(st[1:]) or {}
    if isinstance(st, dict):
        return st
    return {}

def _cell_style_chain(tsc, ri, ci, nrows, ncols):
    """按 pptd.md 表格样式优先级取有效 CellStyle（低→高：cellStyle → bodyStyles → 列类别/行类别）。"""
    if not isinstance(tsc, dict):
        return {}
    chain = []
    if tsc.get("cellStyle"):
        chain.append(tsc["cellStyle"])
    body = tsc.get("bodyStyles")
    if body and 0 < ri < nrows - 1:
        chain.append(body[(ri - 1) % len(body)])
    cols, rows = [], []
    if ci == 0 and tsc.get("firstColumnStyle"):
        cols.append(tsc["firstColumnStyle"])
    if ci == ncols - 1 and tsc.get("lastColumnStyle"):
        cols.append(tsc["lastColumnStyle"])
    if ri == 0 and tsc.get("firstRowStyle"):
        rows.append(tsc["firstRowStyle"])
    if ri == nrows - 1 and tsc.get("lastRowStyle"):
        rows.append(tsc["lastRowStyle"])
    chain += (cols + rows) if tsc.get("rowOverColumn", True) else (rows + cols)
    merged = {}
    for st in chain:
        if isinstance(st, dict):
            merged.update(st)
    return merged

def _set_cell_border(tcell, spec, theme_colors):
    """BorderSpec → 四边 XML 线（单对象=四边同；2 元素=[上下,左右]；4 元素=[上,右,下,左]）。"""
    if isinstance(spec, dict):
        sides = dict.fromkeys(("a:lnT", "a:lnR", "a:lnB", "a:lnL"), spec)
    elif isinstance(spec, (list, tuple)):
        if len(spec) == 2:
            sides = {"a:lnT": spec[0], "a:lnB": spec[0], "a:lnL": spec[1], "a:lnR": spec[1]}
        elif len(spec) == 4:
            sides = {"a:lnT": spec[0], "a:lnR": spec[1], "a:lnB": spec[2], "a:lnL": spec[3]}
        else:
            return
    else:
        return
    tcPr = tcell._tc.get_or_add_tcPr()
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    for tag in order:
        border = sides.get(tag)
        ln = tcPr.find(qn(tag))
        if not isinstance(border, dict):
            if ln is not None:
                tcPr.remove(ln)
            continue
        if ln is None:
            ln = tcPr.makeelement(qn(tag), {})
            anchor = None
            for prev in reversed(order[:order.index(tag)]):
                anchor = tcPr.find(qn(prev))
                if anchor is not None:
                    break
            if anchor is not None:
                anchor.addnext(ln)
            else:
                tcPr.insert(0, ln)
        ln.set("w", str(int(round(float(border.get("width", 1)) * PX))))
        for child in list(ln):
            ln.remove(child)
        if border.get("style") == "none":
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        else:
            sf = ln.makeelement(qn("a:solidFill"), {})
            sf.append(sf.makeelement(qn("a:srgbClr"),
                                     {"val": resolve_color(border.get("color"), theme_colors, "000000")}))
            ln.append(sf)

def _fill_cell(tcell, cell, style, theme, theme_colors, table_id):
    """应用单元格样式链与内联字段（优先级：cell 内联 > cell.textStyle > 表样式链）。"""
    style = style if isinstance(style, dict) else {}
    for src in (style, cell):
        for f in _UNSUPPORTED_CELL_FIELDS:
            if isinstance(src, dict) and src.get(f) is not None:
                _note("table", table_id, f"cell style field '{f}' not supported in local track")
    tstyle = {}
    ts_ref = cell.get("textStyle")
    if isinstance(ts_ref, str) and ts_ref.startswith("$"):
        tstyle = (theme.get("textStyles") or {}).get(ts_ref[1:]) or {}
    cfill = _first(cell.get("fill"), style.get("fill"))
    if cfill:
        typ = cfill.get("type", "solid")
        if typ == "solid":
            tcell.fill.solid()
            tcell.fill.fore_color.rgb = rgb(resolve_color(cfill.get("color"), theme_colors, "FFFFFF"))
        else:
            _note("table", table_id, f"cell fill type '{typ}' not supported in local track")
    border = _first(cell.get("border"), style.get("border"))
    if border is not None:
        _set_cell_border(tcell, border, theme_colors)
    text = cell.get("text", "")
    if not text:
        return
    tf = tcell.text_frame
    tf.word_wrap = True
    txt = html.unescape(re.sub(r"<[^>]+>", "", text))
    lines = [l for l in txt.split("\n") if l.strip()]
    if not lines:
        lines = [txt.strip()] if txt.strip() else []
    if not lines:
        return
    color = resolve_color(
        _first(cell.get("color"), tstyle.get("color"), style.get("color")), theme_colors, "222222")
    size = _first(cell.get("fontSize"), tstyle.get("fontSize"), style.get("fontSize"), 14)
    bold = bool(_first(cell.get("bold"), tstyle.get("bold"), style.get("bold"), False))
    al = _first(cell.get("align"), style.get("align"))
    ff = _first(cell.get("fontFamily"), tstyle.get("fontFamily"), style.get("fontFamily"), "Microsoft YaHei")
    if isinstance(ff, dict):
        latin_f = ff.get("latin", "Arial"); ea_f = ff.get("ea", latin_f)
    else:
        latin_f = ea_f = ff
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(al, list) and al:
            p.alignment = _ALIGN.get(al[0], PP_ALIGN.LEFT)
        r = p.add_run(); r.text = line
        f = r.font
        f.size = Pt(size); f.color.rgb = rgb(color); f.bold = bold
        set_rpr_fonts(r, latin_f, ea_f)
    try:
        tcell.vertical_anchor = MSO_ANCHOR.MIDDLE
    except Exception:
        pass

def add_table(slide, el, theme, theme_colors):
    rows_data = el.get("rows") or []
    colw = el.get("columnWidths") or [1.0] * max(1, len(rows_data[0] if rows_data else [None]))
    ncols = len(colw)
    nrows = len(rows_data)
    if nrows == 0 or ncols == 0:
        return
    x, y, w, h = el.get("bounds", [0, 0, 100, 40])
    table_id = el.get("elementId", "?")
    tsc = _resolve_table_style(el, theme)
    gfx = slide.shapes.add_table(nrows, ncols, to_emu(x), to_emu(y), to_emu(w), to_emu(h))
    table = gfx.table
    table.first_row = False
    table.horz_banding = False
    tw = sum(colw) or ncols
    for ci in range(ncols):
        table.columns[ci].width = to_emu(w * (colw[ci] / tw if tw else w / ncols))
    rowh = el.get("rowHeights") or [1.0] * nrows
    th = sum(rowh) or nrows
    for ri in range(nrows):
        table.rows[ri].height = to_emu(h * (rowh[ri] / th if th else h / nrows))
    grid = [[None] * ncols for _ in range(nrows)]
    for ri, row in enumerate(rows_data):
        ci = 0
        for cell in (row or []):
            while ci < ncols and grid[ri][ci] is not None:
                ci += 1
            if ci >= ncols:
                break
            rs = int(cell.get("rowSpan", 1) or 1)
            cs = int(cell.get("colSpan", 1) or 1)
            try:
                tcell = table.cell(ri, ci)
                if rs > 1 or cs > 1:
                    tcell.merge(table.cell(min(ri + rs - 1, nrows - 1), min(ci + cs - 1, ncols - 1)))
            except Exception:
                tcell = table.cell(ri, ci)
            for rr in range(ri, min(ri + rs, nrows)):
                for cc in range(ci, min(ci + cs, ncols)):
                    if (rr, cc) != (ri, ci):
                        grid[rr][cc] = "merged"
            _fill_cell(tcell, cell, _cell_style_chain(tsc, ri, ci, nrows, ncols), theme, theme_colors, table_id)
            grid[ri][ci] = cell
            ci += cs

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
          "justify": PP_ALIGN.JUSTIFY, "distributed": PP_ALIGN.DISTRIBUTE}
_VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

def parse_rich(text, default_font_size, default_color, default_bold, default_font):
    """把 content.text 转成段落列表，每段为 runs[(txt, {bold,color,size,italic,font})]"""
    if not text:
        return []
    paras = []
    # 按 <p>...</p> 或以换行分段；保留 <ul>/<li> 基本处理
    for chunk in re.split(r"<p[^>]*>|</p>", text):
        chunk = chunk.strip()
        if chunk == "":
            continue
        if re.match(r"^</?[a-zA-Z]+", chunk) is None:
            # 纯文本（可能含换行）→ 一行一段
            for line in chunk.split("\n"):
                line = line.strip()
                if line:
                    paras.append([(html.unescape(line), {})])
            continue
        paras += _parse_html_para(chunk, default_font_size, default_color, default_bold, default_font)
    # 若完全没有 <p>（纯文本多行），上面已处理
    return paras

def _parse_html_para(htmltext, dfs, dc, db, df):
    runs = []
    pos = 0
    pattern = re.compile(r"<(\w+)([^>]*)>|</(\w+)>|(\n)")
    tag_stack = []
    buf = ""
    def flush():
        nonlocal buf
        if buf:
            runs.append((html.unescape(buf), dict(tag_stack)))
            buf = ""
    for m in pattern.finditer(htmltext):
        if m.start() > pos:
            buf += htmltext[pos:m.start()]
        if m.group(4):
            flush()
            pos = m.end()
            continue
        if m.group(1):
            tag = m.group(1).lower()
            attrs = m.group(2) or ""
            if tag in ("strong", "b"):
                tag_stack.append("bold")
                buf += ""
            elif tag in ("em", "i"):
                tag_stack.append("italic")
                buf += ""
            elif tag == "span":
                st = re.search(r'style="([^"]*)"', attrs)
                if st:
                    tag_stack.append(("spanstyle", st.group(1)))
                else:
                    tag_stack.append(("span", ""))
                buf += ""
            elif tag == "br":
                flush()
                buf += "\n"
            else:
                buf += " "  # 其它标签当分隔
        elif m.group(3):
            tag = m.group(3).lower()
            # 关掉对应的栈条目
            for i in range(len(tag_stack) - 1, -1, -1):
                it = tag_stack[i]
                if isinstance(it, str) and it == tag or (isinstance(it, tuple) and it[0] == "span"):
                    del tag_stack[i:]
                    break
        pos = m.end()
    if pos < len(htmltext):
        buf += htmltext[pos:]
    flush()
    # 转成规范 runs
    out = []
    for t, st in runs:
        st2 = dict(st)
        bold = "bold" in st2 or db
        italic = "italic" in st2
        color = dc
        size = dfs
        font = df
        span_style = None
        for k in st2:
            if isinstance(k, tuple) and k[0] == "spanstyle":
                span_style = k[1]
        if span_style:
            cs = re.search(r"color\s*:\s*([#$\w]+)", span_style)
            if cs:
                color = cs.group(1)
            fs = re.search(r"font-size\s*:\s*(\d+)px", span_style)
            if fs:
                size = int(fs.group(1))
        out.append((t, {"bold": bold, "italic": italic, "color": color, "size": size, "font": font}))
    return [out] if out else []

def add_text(cell_or_tf, content, theme):
    # content: {text, style, color, fontSize, fontFamily, bold, italic, lineHeight, align}
    tc = theme.get("textStyles", {})
    st_ref = content.get("style")
    base = {}
    if isinstance(st_ref, str) and st_ref.startswith("$"):
        base = tc.get(st_ref[1:], {}) or {}
    font_size = content.get("fontSize") or base.get("fontSize") or 18
    color_src = content.get("color") or base.get("color") or "$text"
    color = resolve_color(color_src, theme.get("colors", {}), "222222")
    bold = content.get("bold", base.get("bold", False))
    italic = content.get("italic", base.get("italic", False))
    lh = content.get("lineHeight", base.get("lineHeight", 1))
    align = content.get("align")
    fname = content.get("fontFamily") or base.get("fontFamily") or "Microsoft YaHei"
    if isinstance(fname, dict):
        latin_font = fname.get("latin", "Arial")
        ea_font = fname.get("ea", latin_font)
    else:
        latin_font = fname
        ea_font = fname
    tf = cell_or_tf
    tf.word_wrap = True
    paras = parse_rich(content.get("text", ""), font_size, color, bold, latin_font)
    if not paras:
        paras = [[("", {})]]
    for pi, para in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        if lh:
            p.line_spacing = float(lh)
        if align:
            p.alignment = _ALIGN.get(align[0] if isinstance(align, list) else align, PP_ALIGN.LEFT)
        for txt, sty in para:
            run = p.add_run()
            run.text = txt
            f = run.font
            set_rpr_fonts(run, latin_font, ea_font)
            f.size = Pt(sty.get("size", font_size))
            f.bold = bool(sty.get("bold", bold))
            f.italic = bool(sty.get("italic", italic))
            c = resolve_color(sty.get("color", color), theme.get("colors", {}), "222222")
            f.color.rgb = rgb(c)

def main():
    DROPPED.clear()
    _NOTE_SEEN.clear()
    ap = argparse.ArgumentParser()
    ap.add_argument("manifest")
    ap.add_argument("-o", "--output")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    base = os.path.dirname(os.path.abspath(args.manifest))
    m = yaml.safe_load(open(args.manifest, encoding="utf-8"))
    if m.get("version") != "v2":
        sys.exit("仅支持 PPTD v2")
    size = m.get("size", [960, 540])
    theme = m.get("theme", {}) or {}
    theme_colors = theme.get("colors", {}) or {}
    prs = Presentation()
    prs.slide_width = to_emu(size[0])
    prs.slide_height = to_emu(size[1])
    blank = prs.slide_layouts[6]
    pages = m.get("pages") or []
    if not pages:
        sys.exit("pages 列表必须非空：deck 至少要有一页")
    for pi, rel in enumerate(pages):
        page = yaml.safe_load(open(os.path.join(base, rel), encoding="utf-8"))
        if page.get("notes"):
            _note("page", str(pi), "speaker notes not supported in local track")
        if page.get("animations"):
            _note("page", str(pi), "animations not supported in local track")
        slide = prs.slides.add_slide(blank)
        # 背景
        bkg = page.get("background", {}) or {}
        bt = bkg.get("type", "solid")
        bc = resolve_color(bkg.get("color"), theme_colors, "FFFFFF")
        # python-pptx 背景：填一张全屏矩形在底层
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0),
                                    prs.slide_width, prs.slide_height)
        bg.line.fill.background()
        if bt == "solid":
            bg.fill.solid(); bg.fill.fore_color.rgb = rgb(bc)
        elif bt == "gradient":
            stops = bkg.get("stops") or []
            bg.fill.solid()
            bg.fill.fore_color.rgb = rgb(resolve_color(stops[0].get("color") if stops else None, theme_colors, "FFFFFF"))
        else:
            bg.fill.background()
        # 元素（按顺序 = 层序，后加的在上）
        for el in page.get("elements", []):
            et = el.get("elementType")
            eid = el.get("elementId", "?")
            bx = el.get("bounds", [0, 0, 100, 40])
            x, y, w, h = bx
            if el.get("opacity") not in (None, 1, 1.0):
                _note(et, eid, "opacity not supported in local track")
            fl = el.get("flip")
            if fl and any(fl):
                _note(et, eid, "flip not supported in local track")
            if el.get("rotation") and et != "shape":
                _note(et, eid, "rotation not supported for this element type in local track")
            if et == "shape":
                sn = el.get("shapeName", "rect")
                if sn not in SHAPE_MAP:
                    _note(et, eid, f'unknown shapeName "{sn}" rendered as rect')
                clip = find_shape(sn)
                sh = slide.shapes.add_shape(clip, to_emu(x), to_emu(y), to_emu(w), to_emu(h))
                if el.get("rotation"):
                    sh.rotation = float(el["rotation"])
                apply_fill(sh, el.get("fill"), theme_colors, RGBColor(0xD9, 0xD9, 0xD9), owner=(et, eid))
                apply_border(sh, el.get("border"), theme_colors)
                if sn == "custom":
                    sh.fill.background()
                    sh.line.color.rgb = rgb("888888")
            elif et == "text":
                tb = slide.shapes.add_textbox(to_emu(x), to_emu(y), to_emu(w), to_emu(h))
                tf = tb.text_frame
                if isinstance(el.get("content", {}).get("align"), list) and len(el["content"]["align"]) == 2:
                    tf.vertical_anchor = _VALIGN.get(el["content"]["align"][1], MSO_ANCHOR.TOP)
                add_text(tf, el.get("content", {}), theme)
            elif et == "line":
                pts = (el.get("points") or "0,0 1,1").split()
                p0 = list(map(float, pts[0].split(",")))
                p1 = list(map(float, pts[-1].split(",")))
                vb = el.get("viewBox")
                if isinstance(vb, (list, tuple)) and len(vb) == 2 and float(vb[0]) and float(vb[1]):
                    x0b, y0b = x + p0[0] / float(vb[0]) * w, y + p0[1] / float(vb[1]) * h
                    x1b, y1b = x + p1[0] / float(vb[0]) * w, y + p1[1] / float(vb[1]) * h
                else:
                    x0b, y0b, x1b, y1b = x + p0[0], y + p0[1], x + p1[0], y + p1[1]
                if len(pts) > 2:
                    _note(et, eid, "control points degraded to straight segment")
                conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, to_emu(x0b), to_emu(y0b),
                                                  to_emu(x1b), to_emu(y1b))
                bc2 = resolve_color((el.get("border") or {}).get("color"), theme_colors, "000000")
                conn.line.color.rgb = rgb(bc2)
                conn.line.width = Emu(int(round(float((el.get("border") or {}).get("width", 1)) * PX)))
            elif et == "image":
                src = el.get("src")
                path = os.path.join(base, src) if src else None
                if path and os.path.isfile(path):
                    fit = el.get("fit")
                    mode = fit.get("mode") if isinstance(fit, dict) else None
                    mode = mode or "cover"
                    if mode not in ("cover", "contain", "fill"):
                        _note(et, eid, f'unknown fit mode "{mode}" rendered as cover')
                        mode = "cover"
                    if el.get("cropShape"):
                        _note(et, eid, "cropShape not supported in local track")
                    crop = el.get("crop")
                    fr = _crop_fractions(crop)
                    if fr is None:
                        if crop:
                            _note(et, eid, "crop outset not supported in local track; crop ignored")
                        fr = (0.0, 0.0, 0.0, 0.0)
                    with PILImage.open(path) as im:
                        iw, ih = im.size
                    (bl, bt, br, bb), dw, dh = _fit_rect(iw, ih, w, h, fr[0], fr[1], fr[2], fr[3], mode)
                    pic = slide.shapes.add_picture(path, to_emu(x + (w - dw) / 2.0), to_emu(y + (h - dh) / 2.0),
                                                   to_emu(dw), to_emu(dh))
                    _set_src_rect(pic, bl, bt, br, bb)
                else:
                    pass  # 远程/缺失图跳过（SKILL.md 声明的降级）
            elif et == "icon":
                _note(et, eid, "icon elements are not rendered in local track")
            elif et == "table":
                add_table(slide, el, theme, theme_colors)
            elif et == "chart":
                _note(et, eid, "chart elements are not rendered in local track")
            else:
                _note(et, eid, "unknown elementType not rendered in local track")
    out = args.output or os.path.join(base, os.path.splitext(os.path.basename(args.manifest))[0] + ".pptx")
    if os.path.exists(out) and not args.force:
        sys.exit(f"输出已存在（用 --force 覆盖）：{out}")
    prs.save(out)
    print("OK ->", os.path.abspath(out))
    print("slides:", len(prs.slides._sldIdLst))
    for tp, eid, reason in DROPPED:
        print(f"dropped: {tp}/{eid}: {reason}", file=sys.stderr)
    print(f"dropped total: {len(DROPPED)}", file=sys.stderr)

if __name__ == "__main__":
    main()
